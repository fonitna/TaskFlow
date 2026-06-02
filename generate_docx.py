"""
Generate TaskFlow User Manual as a professional DOCX file.
Run: python generate_docx.py
Output: TaskFlow_User_Manual.docx
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── Brand colours ────────────────────────────────────────────────────────────
BLUE        = RGBColor(0x37, 0x8A, 0xDD)   # #378ADD
DARK_BLUE   = RGBColor(0x18, 0x5F, 0xA5)   # #185FA5
LIGHT_BLUE  = RGBColor(0xB5, 0xD4, 0xF4)   # #B5D4F4 (table header fill)
RED         = RGBColor(0xE2, 0x4B, 0x4A)   # #E24B4A
GREEN       = RGBColor(0x63, 0x99, 0x22)   # #639922
AMBER       = RGBColor(0xBA, 0x75, 0x17)   # #BA7517
DARK_GRAY   = RGBColor(0x2C, 0x2C, 0x2A)   # near-black body text
MED_GRAY    = RGBColor(0x60, 0x60, 0x60)   # subtitle / caption
LIGHT_GRAY  = RGBColor(0xF5, 0xF5, 0xF5)   # alternate row fill
NOTE_BG     = RGBColor(0xEF, 0xF6, 0xFF)   # note box background
NOTE_BORDER = RGBColor(0x37, 0x8A, 0xDD)   # note box left border
WARN_BG     = RGBColor(0xFF, 0xF8, 0xEE)
WARN_BORDER = RGBColor(0xBA, 0x75, 0x17)


# ── Helpers ───────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  f'{hex_color[0]:02X}{hex_color[1]:02X}{hex_color[2]:02X}')
    tcPr.append(shd)


def add_page_number(paragraph):
    """Insert 'Page X of Y' into the given paragraph."""
    run = paragraph.add_run('Page ')
    fld = OxmlElement('w:fldChar')
    fld.set(qn('w:fldCharType'), 'begin')
    run._r.append(fld)

    ins = OxmlElement('w:instrText')
    ins.set(qn('xml:space'), 'preserve')
    ins.text = 'PAGE'
    run._r.append(ins)

    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld2)

    paragraph.add_run(' of ')

    run2 = paragraph.add_run()
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'begin')
    run2._r.append(fld3)

    ins2 = OxmlElement('w:instrText')
    ins2.set(qn('xml:space'), 'preserve')
    ins2.text = 'NUMPAGES'
    run2._r.append(ins2)

    fld4 = OxmlElement('w:fldChar')
    fld4.set(qn('w:fldCharType'), 'end')
    run2._r.append(fld4)


def set_paragraph_border_left(paragraph, color: RGBColor, width_pts=18):
    """Add a thick left border to simulate a note box."""
    pPr  = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'),   'single')
    left.set(qn('w:sz'),    str(width_pts))
    left.set(qn('w:space'), '6')
    left.set(qn('w:color'), f'{color[0]:02X}{color[1]:02X}{color[2]:02X}')
    pBdr.append(left)
    pPr.append(pBdr)


def shade_paragraph(paragraph, color: RGBColor):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  f'{color[0]:02X}{color[1]:02X}{color[2]:02X}')
    pPr.append(shd)


def add_horizontal_rule(doc):
    p   = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '6')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), 'CCCCCC')
    pBdr.append(bot)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    return p


def add_note(doc, text: str, kind='info'):
    """Render a > blockquote-style note box."""
    border_color = NOTE_BORDER if kind == 'info' else (RED if kind == 'warning' else WARN_BORDER)
    bg_color     = NOTE_BG     if kind == 'info' else (RGBColor(0xFF,0xEE,0xEE) if kind == 'warning' else WARN_BG)

    p = doc.add_paragraph()
    p.paragraph_format.left_indent    = Cm(0.4)
    p.paragraph_format.right_indent   = Cm(0.4)
    p.paragraph_format.space_before   = Pt(6)
    p.paragraph_format.space_after    = Pt(6)
    set_paragraph_border_left(p, border_color, width_pts=24)
    shade_paragraph(p, bg_color)

    label_map = {'info': 'ℹ Note: ', 'warning': '⚠ Warning: ', 'tip': '💡 Tip: '}
    label = label_map.get(kind, 'Note: ')

    run = p.add_run(label)
    run.bold = True
    run.font.size  = Pt(9.5)
    run.font.color.rgb = border_color

    run2 = p.add_run(text)
    run2.font.size  = Pt(9.5)
    run2.font.color.rgb = DARK_GRAY
    return p


def add_code_inline(paragraph, text: str):
    run = paragraph.add_run(text)
    run.font.name  = 'Courier New'
    run.font.size  = Pt(9)
    run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
    return run


def heading1(doc, text: str):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        run.font.color.rgb = BLUE
        run.font.size      = Pt(18)
        run.font.bold      = True
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    return p


def heading2(doc, text: str):
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        run.font.color.rgb = DARK_BLUE
        run.font.size      = Pt(13)
        run.font.bold      = True
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    return p


def heading3(doc, text: str):
    p = doc.add_heading(text, level=3)
    for run in p.runs:
        run.font.color.rgb = DARK_GRAY
        run.font.size      = Pt(11)
        run.font.bold      = True
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(3)
    return p


def body(doc, text: str, bold_parts: list = None):
    """Add a body paragraph. bold_parts is a list of substrings to bold."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(5)
    p.style = doc.styles['Normal']

    if not bold_parts:
        run = p.add_run(text)
        run.font.size      = Pt(10.5)
        run.font.color.rgb = DARK_GRAY
    else:
        remaining = text
        for bold_word in bold_parts:
            idx = remaining.find(bold_word)
            if idx == -1:
                continue
            before = remaining[:idx]
            if before:
                r = p.add_run(before)
                r.font.size = Pt(10.5)
                r.font.color.rgb = DARK_GRAY
            rb = p.add_run(bold_word)
            rb.bold = True
            rb.font.size = Pt(10.5)
            rb.font.color.rgb = DARK_GRAY
            remaining = remaining[idx+len(bold_word):]
        if remaining:
            r = p.add_run(remaining)
            r.font.size = Pt(10.5)
            r.font.color.rgb = DARK_GRAY
    return p


def bullet(doc, text: str, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent  = Cm(0.5 + level * 0.5)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.font.size      = Pt(10.5)
    run.font.color.rgb = DARK_GRAY
    return p


def numbered(doc, text: str, num: int):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.font.size      = Pt(10.5)
    run.font.color.rgb = DARK_GRAY
    return p


def make_table(doc, headers: list, rows: list, col_widths: list = None):
    """Create a styled table with blue header row and alternating rows."""
    n_cols = len(headers)
    tbl = doc.add_table(rows=1 + len(rows), cols=n_cols)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ''
        set_cell_bg(hdr_cells[i], BLUE)
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.bold           = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size      = Pt(9.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Data rows
    for r_idx, row_data in enumerate(rows):
        cells = tbl.rows[r_idx + 1].cells
        fill  = LIGHT_GRAY if r_idx % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF)
        for c_idx, cell_text in enumerate(row_data):
            cells[c_idx].text = ''
            set_cell_bg(cells[c_idx], fill)
            p = cells[c_idx].paragraphs[0]
            # detect code spans (backtick)
            if '`' in cell_text:
                parts = cell_text.split('`')
                for pi, part in enumerate(parts):
                    if pi % 2 == 1:
                        add_code_inline(p, part)
                    else:
                        if part:
                            run = p.add_run(part)
                            run.font.size      = Pt(9.5)
                            run.font.color.rgb = DARK_GRAY
            else:
                run = p.add_run(cell_text)
                run.font.size      = Pt(9.5)
                run.font.color.rgb = DARK_GRAY
            cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Column widths
    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)

    doc.add_paragraph()  # spacing after table
    return tbl


# ── Document setup ────────────────────────────────────────────────────────────

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3.0)
    section.right_margin  = Cm(2.5)

# Default font
doc.styles['Normal'].font.name  = 'Calibri'
doc.styles['Normal'].font.size  = Pt(10.5)
doc.styles['Normal'].font.color.rgb = DARK_GRAY

# Footer with page numbers
for section in doc.sections:
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)
    for run in fp.runs:
        run.font.size      = Pt(9)
        run.font.color.rgb = MED_GRAY
    fp.add_run('  |  TaskFlow Enterprise Workspace v1.2').font.color.rgb = MED_GRAY


# ═══════════════════════════════════════════════════════════════════════════════
# COVER PAGE
# ═══════════════════════════════════════════════════════════════════════════════

# Blue banner block
cover_banner = doc.add_paragraph()
cover_banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover_banner.paragraph_format.space_before = Pt(60)
cover_banner.paragraph_format.space_after  = Pt(0)
shade_paragraph(cover_banner, BLUE)
r = cover_banner.add_run('  TASKFLOW  ')
r.font.size      = Pt(42)
r.font.bold      = True
r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Subtitle bar
sub_bar = doc.add_paragraph()
sub_bar.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_bar.paragraph_format.space_before = Pt(0)
sub_bar.paragraph_format.space_after  = Pt(0)
shade_paragraph(sub_bar, DARK_BLUE)
rs = sub_bar.add_run('  Enterprise Team Task Management Workspace  ')
rs.font.size      = Pt(13)
rs.font.bold      = False
rs.font.color.rgb = LIGHT_BLUE

doc.add_paragraph()

# Document type label
doc_type = doc.add_paragraph()
doc_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc_type.paragraph_format.space_before = Pt(20)
r2 = doc_type.add_run('USER MANUAL')
r2.font.size      = Pt(22)
r2.font.bold      = True
r2.font.color.rgb = DARK_GRAY

# Version / date row
ver_p = doc.add_paragraph()
ver_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
ver_p.paragraph_format.space_before = Pt(8)
rv = ver_p.add_run('Version 1.2  |  June 2026')
rv.font.size      = Pt(11)
rv.font.color.rgb = MED_GRAY

doc.add_paragraph()
doc.add_paragraph()

# Decorative separator
sep = doc.add_paragraph()
sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs2 = sep.add_run('━' * 32)
rs2.font.color.rgb = LIGHT_BLUE
rs2.font.size      = Pt(12)

doc.add_paragraph()

# Summary blurb
blurb = doc.add_paragraph()
blurb.alignment = WD_ALIGN_PARAGRAPH.CENTER
blurb.paragraph_format.space_before = Pt(8)
rb = blurb.add_run(
    'A full-stack, responsive project and task management application\n'
    'with Kanban boards, JWT authentication, and a persistent SQLite backend.'
)
rb.font.size      = Pt(10.5)
rb.font.color.rgb = MED_GRAY

doc.add_paragraph()
doc.add_paragraph()

# Bottom info table on cover
info_tbl = doc.add_table(rows=3, cols=2)
info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
info_data = [
    ('Access URL (Dev)',    'http://localhost:3000'),
    ('Access URL (Docker)', 'http://localhost:3500'),
    ('Default Password',    'password123'),
]
for i, (label, value) in enumerate(info_data):
    cells = info_tbl.rows[i].cells
    cells[0].text = ''
    set_cell_bg(cells[0], BLUE)
    r_l = cells[0].paragraphs[0].add_run(label)
    r_l.bold           = True
    r_l.font.size      = Pt(9.5)
    r_l.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    cells[1].text = ''
    set_cell_bg(cells[1], LIGHT_GRAY)
    r_v = cells[1].paragraphs[0].add_run(value)
    r_v.font.size      = Pt(9.5)
    r_v.font.color.rgb = DARK_GRAY

for row in info_tbl.rows:
    row.cells[0].width = Cm(5)
    row.cells[1].width = Cm(8)

# Page break after cover
doc.add_page_break()


# ═══════════════════════════════════════════════════════════════════════════════
# TABLE OF CONTENTS PAGE
# ═══════════════════════════════════════════════════════════════════════════════

toc_title = doc.add_paragraph()
toc_title.paragraph_format.space_before = Pt(0)
shade_paragraph(toc_title, BLUE)
rt = toc_title.add_run('  Table of Contents')
rt.font.size      = Pt(16)
rt.font.bold      = True
rt.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

doc.add_paragraph()

toc_entries = [
    ('1.', 'Introduction'),
    ('2.', 'Accessing TaskFlow'),
    ('3.', 'Authentication'),
    ('   3.1', 'Sign In'),
    ('   3.2', 'Register a New Account'),
    ('   3.3', 'Sign Out'),
    ('4.', 'Navigation'),
    ('   4.1', 'Sidebar'),
    ('   4.2', 'Header Bar'),
    ('   4.3', 'Mobile Navigation'),
    ('5.', 'Dashboard'),
    ('6.', 'Projects'),
    ('   6.1', 'Viewing All Projects'),
    ('   6.2', 'Creating a Project'),
    ('   6.3', 'Reordering Projects'),
    ('7.', 'Project Board'),
    ('   7.1', 'Kanban Board View'),
    ('   7.2', 'List View'),
    ('   7.3', 'Filtering Tasks'),
    ('8.', 'Tasks'),
    ('   8.1', 'Creating a Task'),
    ('   8.2', 'Viewing Task Details'),
    ('   8.3', 'Editing a Task'),
    ('   8.4', 'Moving Tasks Between Columns'),
    ('   8.5', 'Adding and Removing Labels'),
    ('   8.6', 'Deleting a Task'),
    ('   8.7', 'Bulk Operations'),
    ('9.', 'Comments'),
    ('10.', 'Search'),
    ('   10.1', 'Quick Search Modal'),
    ('   10.2', 'Full Search Page'),
    ('11.', 'Profile'),
    ('12.', 'Settings'),
    ('   12.1', 'Theme Mode'),
    ('   12.2', 'Project Settings'),
    ('13.', 'Keyboard Shortcuts'),
    ('14.', 'Mobile Usage'),
    ('15.', 'Default Accounts'),
]

for num, title in toc_entries:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    is_top = not num.startswith(' ')
    indent = Cm(0) if is_top else Cm(0.8)
    p.paragraph_format.left_indent = indent

    rn = p.add_run(f'{num.strip()}  ')
    rn.bold           = is_top
    rn.font.size      = Pt(10.5 if is_top else 10)
    rn.font.color.rgb = BLUE if is_top else MED_GRAY

    rt2 = p.add_run(title)
    rt2.bold           = is_top
    rt2.font.size      = Pt(10.5 if is_top else 10)
    rt2.font.color.rgb = DARK_GRAY if is_top else MED_GRAY

doc.add_page_break()


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 1 — INTRODUCTION
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '1. Introduction')
body(doc,
    'TaskFlow is a team task management workspace that organizes work across multiple projects '
    'using a Kanban board model. It supports multiple users, project-based task tracking, '
    'real-time collaboration through comments, and role-based team member assignment.'
)
body(doc, 'Key concepts used throughout this manual:')

make_table(doc,
    headers=['Concept', 'Description'],
    rows=[
        ['Project',  'A named workspace (e.g. "Website Redesign") containing tasks and assigned team members'],
        ['Task',     'A unit of work with a title, status, priority, assignee, due date, and labels'],
        ['Status',   'One of four stages: To Do → In Progress → In Review → Done'],
        ['Priority', 'One of four levels: Low, Normal, High, Urgent'],
        ['Label',    'A free-text tag attached to a task for categorization (e.g. #frontend, #bug)'],
    ],
    col_widths=[4, 12]
)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 2 — ACCESSING
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '2. Accessing TaskFlow')
make_table(doc,
    headers=['Mode', 'URL'],
    rows=[
        ['Local development', 'http://localhost:3000'],
        ['Docker deployment', 'http://localhost:3500'],
    ],
    col_widths=[6, 10]
)
body(doc,
    'Open the URL in any modern browser. The application loads with a full-page spinner while it '
    'restores your session. If no session is found, you are redirected to the login page.'
)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 3 — AUTHENTICATION
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '3. Authentication')

heading2(doc, '3.1 Sign In')
body(doc, 'The Sign In tab is shown by default on the login page.')
body(doc, 'Steps:')
numbered(doc, 'The page displays all registered users as selectable profile cards.', 1)
numbered(doc, 'Click the card of the user you want to sign in as. A blue ring appears and a checkmark replaces the initials.', 2)
numbered(doc, 'Your email address appears below the card list.', 3)
numbered(doc, 'Enter your password in the password field. Click the eye icon to show or hide the password.', 4)
numbered(doc, 'Click Proceed to Dashboard.', 5)
add_note(doc, 'Default credentials for seeded accounts: password is password123 for all pre-loaded users. See Section 15 for the full list.', kind='info')

heading2(doc, '3.2 Register a New Account')
numbered(doc, 'Click the Join Team tab at the top of the login page.', 1)
numbered(doc, 'Fill in the registration form:', 2)
make_table(doc,
    headers=['Field', 'Description', 'Validation'],
    rows=[
        ['Professional Name', 'Your display name', 'Minimum 2 characters, max 25'],
        ['Email Address',     'Unique email address', 'Must contain @'],
        ['Password',          'Login password', 'Minimum 6 characters'],
        ['Role',              'Select from dropdown', 'Product Manager, Developer, Designer, QA Specialist, Content Writer, Team Lead'],
        ['Avatar Color',      'Click one of 7 color swatches', 'Required (blue selected by default)'],
    ],
    col_widths=[4, 6.5, 5.5]
)
numbered(doc, 'A live preview card appears below as you type your name, showing how your avatar will look.', 3)
numbered(doc, 'Click Complete Onboarding & Enter.', 4)
numbered(doc, 'On success, you are automatically logged in and taken to the Dashboard.', 5)
add_note(doc, 'Email addresses must be unique. If the address is already registered, an error banner appears.', kind='info')

heading2(doc, '3.3 Sign Out')
numbered(doc, 'Click your user avatar in the top-right corner of the header bar.', 1)
numbered(doc, 'A small dropdown appears showing your name.', 2)
numbered(doc, 'Click Sign Out Session.', 3)
body(doc, 'You are returned to the login page and your session token is cleared.')

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 4 — NAVIGATION
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '4. Navigation')

heading2(doc, '4.1 Sidebar')
body(doc, 'The sidebar appears on the left side of the screen on tablets and desktops.')
make_table(doc,
    headers=['Navigation Item', 'Destination'],
    rows=[
        ['TaskFlow logo',    '/ — Dashboard (home page)'],
        ['Dashboard',        '/ — Your personal stats and overdue tasks'],
        ['All Projects',     '/projects — List of all projects'],
        ['Faceted Search',   '/search — Advanced search with filters'],
        ['User Profile',     '/profile — Edit your name, color, and view the team directory'],
        ['Global Settings',  '/settings — Project configuration and theme'],
    ],
    col_widths=[5, 11]
)
body(doc,
    'Projects quick-list: Below the main navigation, the sidebar lists all projects with colored '
    'dots and the count of unresolved (non-done) tasks. Click any project name to open its board directly.'
)
body(doc,
    'Collapsing the sidebar (desktop only): Click the chevron arrow at the top-right of the sidebar '
    'to collapse it to icon-only mode. Click again to expand.'
)

heading2(doc, '4.2 Header Bar')
body(doc, 'The header is always visible at the top of the screen.')
make_table(doc,
    headers=['Element', 'Action'],
    rows=[
        ['Search bar (desktop) / search icon (mobile)', 'Opens the Quick Search modal'],
        ['Moon / Sun icon',  'Toggles dark mode / light mode'],
        ['Bell icon',        'Opens the notifications panel'],
        ['User avatar (top right)', 'Opens the account dropdown (Sign Out)'],
    ],
    col_widths=[7.5, 8.5]
)

heading2(doc, '4.3 Mobile Navigation')
body(doc,
    'On screens narrower than 480 px, the sidebar is replaced by a fixed bottom navigation bar '
    'with five tabs:'
)
make_table(doc,
    headers=['Tab Icon', 'Destination'],
    rows=[
        ['Home',        'Dashboard'],
        ['Grid',        'All Projects'],
        ['Plus (center)', 'Create Task — quick add from any screen'],
        ['Search',      'Faceted Search'],
        ['Person',      'User Profile'],
    ],
    col_widths=[4, 12]
)
body(doc,
    'On screens between 480 px and 768 px, a hamburger menu button appears in the header. '
    'Tap it to slide in the sidebar as an overlay. Tap outside or press the close button to dismiss.'
)

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 5 — DASHBOARD
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '5. Dashboard')
body(doc,
    'The Dashboard is your personal overview screen. It shows only tasks assigned to you.'
)

heading3(doc, 'Welcome Banner')
body(doc, 'At the top of the page, a card shows:')
bullet(doc, 'Your name with a greeting and today\'s date')
bullet(doc, 'A Quick Create Task button (blue, top right) — opens the task creation drawer with status set to "To Do"')

heading3(doc, 'Weekly Progress')
body(doc, 'A bar chart shows the distribution of your tasks across four statuses:')
make_table(doc,
    headers=['Bar', 'Color', 'Meaning'],
    rows=[
        ['To Do',     'Light blue', 'Tasks not yet started'],
        ['Active',    'Amber',      'Tasks currently in progress'],
        ['In Review', 'Blue',       'Tasks awaiting review'],
        ['Done',      'Dark blue',  'Completed tasks'],
    ],
    col_widths=[4, 4, 8]
)
body(doc, 'Below the chart, three summary numbers are displayed:')
bullet(doc, 'Backlogged — count of your To Do tasks')
bullet(doc, 'Overdue — count of tasks past their due date and not done (shown in red)')
bullet(doc, 'Productivity Rate — percentage of your tasks that are Done (shown in green)')

heading3(doc, 'Urgent & Overdue')
body(doc,
    'The red card shows up to 3 of your tasks with a due date in the past and not yet marked Done. '
    'Each entry shows the task title, project name, and due date. Click any task to open its detail modal. '
    'If more than 3 tasks are overdue, a "+ X more overdue tasks" message appears at the bottom.'
)

heading3(doc, 'Active Projects')
body(doc,
    'Shows up to 3 projects with a progress bar indicating what percentage of tasks are Done. '
    'Click a project name to navigate to its board. Click View all project boards to go to the Projects page.'
)

heading3(doc, 'Active Responsibilities')
body(doc, 'A list of your active (non-Done) tasks showing task title, project, status, and priority badge.')
body(doc, 'If you have more than 4 active tasks, a note indicates how many more exist. Click Filter List to go to the full search page.')

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 6 — PROJECTS
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '6. Projects')

heading2(doc, '6.1 Viewing All Projects')
body(doc, 'Navigate to All Projects in the sidebar (or /projects).')
body(doc, 'Each project card displays:')
bullet(doc, 'Project name and ID')
bullet(doc, 'Task counts: Total, Active (in progress), Resolved (done)')
bullet(doc, 'A progress bar showing the percentage of tasks completed')
bullet(doc, 'Team members shown as overlapping avatar circles')
body(doc, 'Click anywhere on a project card to open its board.')

heading2(doc, '6.2 Creating a Project')
numbered(doc, 'Click Create New Project (top right of the Projects page).', 1)
numbered(doc, 'A creation form slides in from the right.', 2)
numbered(doc, 'Fill in the following fields:', 3)
make_table(doc,
    headers=['Field', 'Description', 'Validation'],
    rows=[
        ['Project Name',     'The name of the project', '3–50 characters, must be unique'],
        ['Interface Color',  'Choose one of 7 color swatches', 'Required'],
        ['Assign Members',   'Click member avatars to add them to the project', 'At least 1 member required'],
    ],
    col_widths=[4, 7, 5]
)
numbered(doc, 'Click Create Project to save, or Cancel to discard.', 4)
add_note(doc, 'If the project name is already taken, an error message appears: "Project name must be unique in this workspace."', kind='tip')

heading2(doc, '6.3 Reordering Projects')
numbered(doc, 'Hover over a project card — a grip icon appears.', 1)
numbered(doc, 'Click and hold the card, then drag it to a new position.', 2)
numbered(doc, 'Release to drop. The new order is saved automatically.', 3)
add_note(doc, 'While dragging, the card becomes semi-transparent and surrounding cards shift to indicate the drop position.', kind='info')

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 7 — PROJECT BOARD
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '7. Project Board')
body(doc,
    'Click any project name to open its board. At the top you will see the project name, '
    'creation date, and total task count.'
)

heading2(doc, '7.1 Kanban Board View')
body(doc, 'The default view shows four columns, one per status:')
make_table(doc,
    headers=['Column', 'Status Value', 'Description'],
    rows=[
        ['To Do',       '`todo`',        'Tasks not yet started'],
        ['In Progress', '`in_progress`', 'Tasks currently being worked on'],
        ['In Review',   '`in_review`',   'Tasks awaiting review or approval'],
        ['Done',        '`done`',        'Completed tasks'],
    ],
    col_widths=[4, 4, 8]
)
body(doc, 'Each column shows:')
bullet(doc, 'A colored dot, column name, and task count')
bullet(doc, 'A scrollable list of task cards')
bullet(doc, 'An Add new task button at the bottom of the column')
body(doc, 'Task card contents: title, priority badge, due date (red if overdue), assignee avatar, comment count.')
body(doc, 'Click any task card to open the Task Detail Modal.')

heading2(doc, '7.2 List View')
body(doc, 'Click the List View button (next to Kanban Board) to switch to a table layout.')
make_table(doc,
    headers=['Column', 'Sortable', 'Description'],
    rows=[
        ['Checkbox',    '—',   'Select tasks for bulk operations'],
        ['ID',          '—',   'Task ID (monospace)'],
        ['Task Title',  '✓',   'Title and attached labels'],
        ['Priority',    '✓',   'Color-coded priority badge'],
        ['Status',      '—',   'Inline dropdown to change status immediately'],
        ['Assignee',    '—',   'Avatar and first name'],
        ['Due Date',    '✓',   'Date value or "—" if not set'],
    ],
    col_widths=[4, 3, 9]
)
body(doc, 'Sorting: Click a column header with the ↕ icon to sort ascending. Click again to sort descending.')
body(doc, 'Changing status inline: Click the status dropdown in any row and select a new status. The change is saved immediately.')

heading2(doc, '7.3 Filtering Tasks')
body(doc, 'The filter bar appears below the project header on both Kanban and List views.')
heading3(doc, 'Keyword Search')
body(doc, 'Type in the search box to filter tasks by title or description in real time.')
heading3(doc, 'Assignee Filter')
bullet(doc, 'Click any team member\'s avatar in the filter bar to show only their tasks. Click again to deselect.')
bullet(doc, 'Multiple assignees can be selected simultaneously.')
heading3(doc, 'Priority and Status Filters')
numbered(doc, 'Click the Priority / Status button to expand the filter panel.', 1)
numbered(doc, 'Click priority pills (Low, Normal, High, Urgent) to toggle them on/off.', 2)
numbered(doc, 'Click status pills (To Do, In Progress, In Review, Done) to toggle them on/off.', 3)
numbered(doc, 'Multiple options can be active at the same time.', 4)
heading3(doc, 'Clearing Filters')
body(doc, 'When any filter is active, a red Clear All Filters link appears. Click it to reset all filters at once.')

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 8 — TASKS
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '8. Tasks')

heading2(doc, '8.1 Creating a Task')
body(doc, 'There are three ways to create a task:')
bullet(doc, 'Method 1 — Quick Create Task: Click Quick Create Task on the Dashboard banner. Drawer opens with status "To Do".')
bullet(doc, 'Method 2 — Add Task: On any Project Board, click the blue Add Task button in the top-right area.')
bullet(doc, 'Method 3 — Column button: Click + Add new task at the bottom of any Kanban column. The drawer pre-selects that column\'s status.')
body(doc, 'Task creation drawer fields:')
make_table(doc,
    headers=['Field', 'Required', 'Description'],
    rows=[
        ['Task Title',     '✓',  'Short description of the task. Max 200 characters.'],
        ['Description',    '—',  'Detailed scope or requirements. Resizable textarea.'],
        ['Project',        '✓',  'Select which project this task belongs to.'],
        ['Assignee',       '—',  'Click a team member\'s avatar to assign. Click "—" for unassigned.'],
        ['Priority',       '—',  'Segmented control: Low / Normal / High / Urgent. Default: Normal.'],
        ['Due Date',       '—',  'Date picker. Warning appears if selected date is in the past.'],
        ['Labels / Tags',  '—',  'Type a tag name, press Enter or click Add. Click × to remove.'],
    ],
    col_widths=[3.5, 2.5, 10]
)
body(doc, 'Click Create Task (bottom right of the drawer) to save, or Cancel to discard.')
add_note(doc, 'The drawer closes immediately after creating. The new task appears in the correct Kanban column within seconds.', kind='info')

heading2(doc, '8.2 Viewing Task Details')
body(doc, 'Click any task card (Kanban view) or task title (List view) to open the Task Detail Modal.')
body(doc, 'The modal is divided into two panels:')
body(doc, 'Left panel:')
bullet(doc, 'Task ID and project name (header)')
bullet(doc, 'Task title — click to edit inline')
bullet(doc, 'Description — click to expand and edit')
bullet(doc, 'Labels section')
bullet(doc, 'Comment thread')
body(doc, 'Right panel (Task Metadata):')
bullet(doc, 'Status dropdown')
bullet(doc, 'Assignee dropdown')
bullet(doc, 'Priority dropdown')
bullet(doc, 'Due Date picker')
bullet(doc, 'Created date (read-only)')

heading2(doc, '8.3 Editing a Task')
body(doc, 'All task fields can be edited directly in the Task Detail Modal.')
heading3(doc, 'Editing the Title')
numbered(doc, 'Click the task title text. It transforms into an editable input field with a blue border.', 1)
numbered(doc, 'Type the new title (max 200 characters).', 2)
numbered(doc, 'Press Enter or click outside the field to save. The change is saved automatically.', 3)
heading3(doc, 'Editing the Description')
numbered(doc, 'Click anywhere in the description textarea.', 1)
numbered(doc, 'Type or modify the text.', 2)
numbered(doc, 'Click Save Details to save, or Cancel to discard changes.', 3)
heading3(doc, 'Changing Other Fields')
make_table(doc,
    headers=['Field', 'How to Change'],
    rows=[
        ['Status',    'Click the Status dropdown in the right panel and select a new value. Saved immediately.'],
        ['Priority',  'Click the Priority dropdown and select Low, Normal, High, or Urgent.'],
        ['Assignee',  'Click the Assignee dropdown. Select a member, or Unassigned to clear.'],
        ['Due Date',  'Click the Due Date field, use the date picker. Clear the field to remove the date.'],
    ],
    col_widths=[3.5, 12.5]
)

heading2(doc, '8.4 Moving Tasks Between Columns')
make_table(doc,
    headers=['Method', 'How to Use'],
    rows=[
        ['Drag and drop\n(Kanban view)',   'Click and hold a task card. Drag to the target column (highlighted with blue dashed border). Release to drop. A green ring animation confirms.'],
        ['Status dropdown\n(Detail Modal)', 'Change the Status field in the right panel of the Task Detail Modal.'],
        ['Inline status change\n(List View)', 'Click the status dropdown in the task row and select a new status.'],
    ],
    col_widths=[4.5, 11.5]
)

heading2(doc, '8.5 Adding and Removing Labels')
heading3(doc, 'Adding a Label')
numbered(doc, 'In the Labels section of the left panel, click + Add Label.', 1)
numbered(doc, 'Type the label name (e.g. frontend, bug, design).', 2)
numbered(doc, 'Press Enter to confirm. The label appears as a blue pill prefixed with #.', 3)
heading3(doc, 'Removing a Label')
body(doc, 'Click the × button on any label pill to remove it immediately.')
body(doc, 'Labels also appear as gray tags below the task title in List View.')

heading2(doc, '8.6 Deleting a Task')
numbered(doc, 'Click the red trash icon in the top-right corner of the modal header.', 1)
numbered(doc, 'A confirmation dialog appears: "Are you sure you want to delete this task? This action cannot be undone."', 2)
numbered(doc, 'Click OK to confirm. The task is deleted and the modal closes.', 3)
add_note(doc, 'Deleting a task also deletes all its comments. This action cannot be undone.', kind='warning')

heading2(doc, '8.7 Bulk Operations')
body(doc, 'Bulk operations are only available in List View.')
heading3(doc, 'Selecting Tasks')
bullet(doc, 'Click the checkbox on the left of any task row to select it.')
bullet(doc, 'Click the checkbox in the header row to select or deselect all visible tasks.')
body(doc, 'When one or more tasks are selected, the bulk action bar appears at the top of the table.')
make_table(doc,
    headers=['Action', 'How to Use', 'Description'],
    rows=[
        ['Move to status',   'Select from the dropdown',  'Moves all selected tasks to the chosen status'],
        ['Unassign Users',   'Click the button',           'Removes assignees from all selected tasks'],
        ['Delete',           'Click the red button',       'Deletes all selected tasks (with confirmation dialog)'],
        ['Deselect',         'Click the link',             'Clears the selection without taking action'],
    ],
    col_widths=[4, 4.5, 7.5]
)
add_note(doc, 'Bulk delete shows a confirmation dialog: "Are you sure you want to delete these X tasks?"', kind='warning')

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 9 — COMMENTS
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '9. Comments')
body(doc, 'Every task has a threaded comment section in the lower-left panel of the Task Detail Modal.')

heading3(doc, 'Viewing Comments')
bullet(doc, 'All comments appear in chronological order (oldest first).')
bullet(doc, 'Each comment shows the author\'s avatar, name, comment text, and relative time (e.g. "2h ago", "3d ago").')

heading3(doc, 'Posting a Comment')
numbered(doc, 'Click the text area at the bottom of the comment section. Placeholder: "Write a response... (Press Send)"', 1)
numbered(doc, 'Type your comment. Maximum 2 000 characters.', 2)
numbered(doc, 'A character counter appears when you reach 1 800 characters.', 3)
numbered(doc, 'Click the Send button (paper plane icon).', 4)
numbered(doc, 'A brief spinner animation confirms the comment is being posted. It appears in the list within a second.', 5)
add_note(doc, 'Comments are posted as the currently logged-in user. You cannot post comments when not logged in.', kind='info')

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 10 — SEARCH
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '10. Search')

heading2(doc, '10.1 Quick Search Modal')
body(doc, 'The Quick Search modal provides fast, instant search across all content.')
heading3(doc, 'Opening the Modal')
bullet(doc, 'Press Ctrl+K (Windows/Linux) or Cmd+K (Mac)')
bullet(doc, 'Or click the search bar in the header (desktop) / search icon (mobile)')
heading3(doc, 'Using the Modal')
numbered(doc, 'Start typing your search term. Results update in real time.', 1)
numbered(doc, 'Results are grouped into three sections: Tasks (up to 5), Projects (all matches), Team Members (all matches).', 2)
numbered(doc, 'Click any result to navigate — Task opens the Task Detail Modal, Project opens the board.', 3)
numbered(doc, 'Press Enter or click "Press Enter to view all results page →" to go to the full search page.', 4)
heading3(doc, 'Closing the Modal')
bullet(doc, 'Press Escape')
bullet(doc, 'Click anywhere outside the modal')

heading2(doc, '10.2 Full Search Page')
body(doc, 'Navigate to Faceted Search in the sidebar (or /search).')
heading3(doc, 'Searching')
numbered(doc, 'Type your keyword in the main search bar.', 1)
numbered(doc, 'Press Enter or click the Search button.', 2)
heading3(doc, 'Filtering Results (left panel)')
make_table(doc,
    headers=['Filter', 'Type', 'Options'],
    rows=[
        ['Project',   'Dropdown', 'All Projects, or select a specific project'],
        ['Status',    'Dropdown', 'All Statuses / To Do / In Progress / In Review / Done'],
        ['Priority',  'Dropdown', 'All Priorities / Low / Normal / High / Urgent'],
        ['Assignee',  'Dropdown', 'All Members, or select a specific team member'],
    ],
    col_widths=[4, 3, 9]
)
body(doc, 'All filters are applied together (AND logic). Results update as you change dropdowns.')
body(doc, 'Each result card shows: Task ID, project name, priority badge, task title (click to open detail), description preview, labels, assignee avatar, and status badge.')
body(doc, 'Resetting filters: Click the red Reset link that appears in the filter panel when any filter is active.')

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 11 — PROFILE
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '11. Profile')
body(doc, 'Navigate to User Profile in the sidebar (or /profile).')

heading3(doc, 'Editing Your Profile')
make_table(doc,
    headers=['Field', 'Description'],
    rows=[
        ['Display Name',       'Your visible name across the workspace. Min 2 characters.'],
        ['Avatar Initials',    'Automatically derived from your name when saved.'],
        ['Accent Theme Color', 'Click one of 7 color swatches to change your avatar color.'],
    ],
    col_widths=[5, 11]
)
body(doc, 'The avatar preview at the top of the form updates live as you change the color.')
body(doc, 'Click Save Profile Config to persist changes to the database. A success toast confirms the save.')

heading3(doc, 'Notification Preferences')
body(doc, 'Three checkboxes allow you to configure notification preferences:')
bullet(doc, 'Dispatch email digests on assigned tasks')
bullet(doc, 'Enable native system browser push triggers')
bullet(doc, 'Receive relative digest summaries daily at 8:00 AM')

heading3(doc, 'Team Directory')
body(doc, 'The right card shows a read-only list of all workspace members with their name, role, and user ID. Your own entry is highlighted with a YOU badge in blue.')

heading3(doc, 'Signing Out from Profile Page')
body(doc, 'Click Sign Out Session (bottom left of the form) to log out and return to the login page.')

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 12 — SETTINGS
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '12. Settings')
body(doc, 'Navigate to Global Settings in the sidebar (or /settings).')

heading2(doc, '12.1 Theme Mode')
body(doc, 'At the top of the settings page, a segmented toggle switches between:')
make_table(doc,
    headers=['Option', 'Effect'],
    rows=[
        ['Light', 'Warm beige background (#F1EFE8), dark text — default'],
        ['Dark',  'Deep black background (#121212), light text'],
    ],
    col_widths=[3, 13]
)
body(doc, 'The theme change applies immediately and is remembered for your next visit.')
add_note(doc, 'You can also toggle dark mode from the moon/sun icon in the header bar at any time.', kind='tip')

heading2(doc, '12.2 Project Settings')
heading3(doc, 'Selecting a Project')
numbered(doc, 'Use the Active Selected Project dropdown to choose which project to configure.', 1)
numbered(doc, 'Click Go to project board → to navigate to that project\'s board directly.', 2)
heading3(doc, 'Editing Project Metadata')
make_table(doc,
    headers=['Field', 'Description', 'Validation'],
    rows=[
        ['Project Name',    'Rename the project',                            '3–50 characters, must be unique'],
        ['Accent Color',    'Change the project\'s color identity',           'Select from 7 swatches'],
        ['Collaborators',   'Check/uncheck team members for this project',    'At least 1 must remain checked'],
    ],
    col_widths=[4, 7, 5]
)
body(doc, 'Click Save Project Config to save all changes. A success toast confirms.')
heading3(doc, 'Column Management (right panel)')
bullet(doc, 'Displays the current list of board columns.')
bullet(doc, 'The four default columns (To Do, In Progress, In Review, Done) are locked and cannot be deleted.')
bullet(doc, 'Custom columns can be added and deleted.')
bullet(doc, 'To add: type a column name and click Add.')
bullet(doc, 'To delete: click the trash icon next to a custom column.')
add_note(doc, 'Column management is a preview feature. Changes are shown locally but the board always uses the four built-in statuses.', kind='info')

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 13 — KEYBOARD SHORTCUTS
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '13. Keyboard Shortcuts')
make_table(doc,
    headers=['Shortcut', 'Action'],
    rows=[
        ['Ctrl+K / Cmd+K',          'Open Quick Search modal'],
        ['Escape',                   'Close any open modal, drawer, or search overlay'],
        ['Enter (search modal)',     'Navigate to full search results page'],
        ['Enter (task title edit)',  'Save the title change'],
        ['Enter (label input)',      'Add the typed label to the task'],
    ],
    col_widths=[5.5, 10.5]
)

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 14 — MOBILE USAGE
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '14. Mobile Usage')
body(doc, 'TaskFlow is fully responsive across all screen sizes.')
make_table(doc,
    headers=['Screen Width', 'Layout'],
    rows=[
        ['≥ 768 px (tablet/desktop)',  'Sidebar always visible, collapsible'],
        ['480 px – 767 px',            'Sidebar hidden; hamburger menu in header opens it as an overlay'],
        ['< 480 px',                   'Sidebar replaced by a bottom navigation bar with 5 icon tabs'],
    ],
    col_widths=[5.5, 10.5]
)
body(doc, 'Tips for mobile use:')
bullet(doc, 'Use the center plus button in the bottom nav to quickly create a task from any screen.')
bullet(doc, 'The Kanban board scrolls horizontally on small screens — swipe left and right to see all four columns.')
bullet(doc, 'The Task Detail Modal is scrollable; swipe up on the modal to see the comments section.')
bullet(doc, 'Tap outside an open modal or drawer to close it.')

add_horizontal_rule(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 15 — DEFAULT ACCOUNTS
# ═══════════════════════════════════════════════════════════════════════════════

heading1(doc, '15. Default Accounts')
body(doc, 'The following accounts are pre-loaded in the system. All share the password password123.')
make_table(doc,
    headers=['Name', 'Email', 'Role'],
    rows=[
        ['Alex Chen',     'alex@taskflow.dev',   'PM'],
        ['Priya Sharma',  'priya@taskflow.dev',  'Developer'],
        ['Jordan Lee',    'jordan@taskflow.dev', 'Designer'],
        ['Marco Rossi',   'marco@taskflow.dev',  'Developer'],
        ['Sara Kim',      'sara@taskflow.dev',   'QA'],
    ],
    col_widths=[4.5, 6.5, 5]
)
add_note(doc, 'Change the default passwords after first login in a production environment. Passwords are hashed with bcrypt (cost factor 12) and are never stored in plain text.', kind='warning')


# ═══════════════════════════════════════════════════════════════════════════════
# FINAL PAGE — document info
# ═══════════════════════════════════════════════════════════════════════════════

doc.add_page_break()
closing = doc.add_paragraph()
closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
closing.paragraph_format.space_before = Pt(80)
shade_paragraph(closing, BLUE)
rc = closing.add_run('  TaskFlow Enterprise Workspace v1.2  ')
rc.font.size      = Pt(14)
rc.font.bold      = True
rc.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

cl2 = doc.add_paragraph()
cl2.alignment = WD_ALIGN_PARAGRAPH.CENTER
rc2 = cl2.add_run('User Manual — June 2026')
rc2.font.size      = Pt(10)
rc2.font.color.rgb = MED_GRAY

cl3 = doc.add_paragraph()
cl3.alignment = WD_ALIGN_PARAGRAPH.CENTER
rc3 = cl3.add_run('Licensed under Apache 2.0')
rc3.font.size      = Pt(9)
rc3.font.color.rgb = MED_GRAY


# ── Save ──────────────────────────────────────────────────────────────────────

output_path = 'TaskFlow_User_Manual.docx'
doc.save(output_path)
print(f'[OK] Saved: {output_path}')
